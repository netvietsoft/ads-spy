import sys
sys.stdout.reconfigure(encoding='utf-8')
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print('Reading Superfiliate_Enriched.csv...')
with open(r'G:\AFF\Superfiliate_Enriched.csv', 'r', encoding='utf-8', errors='replace') as f:
    rows = list(csv.reader(f))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Superfiliate Sites (960)'

# Styling
header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Navy
header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')

shopify_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
shopify_font = Font(name='Segoe UI', size=10, color='166534', bold=True)

live_fill = PatternFill(start_color='EFF6FF', end_color='EFF6FF', fill_type='solid')

headers = [
    'Domain', 'Tên công ty / Shop', 'Website', 'Ngành hàng', 'Shopify?',
    'Shop ID', 'DT tháng Local (USD)', 'DT BuiltWith (năm/tháng)', 'Chi tiêu công nghệ (BW)',
    'Link tham gia', 'Ghi chú chi tiết', 'Trạng thái DNS'
]

ws.append(headers)

for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for r_idx, r in enumerate(rows[1:], start=2):
    domain = r[0]
    company = r[1]
    web = r[2]
    vertical = r[3]
    is_shop = r[4] == '1'
    shop_id = r[5]
    rev_month = float(r[6]) if r[6] else None
    bw_rev = r[7]
    bw_spend = r[8]
    join_url = r[9]
    notes = r[10]
    live = r[11] == '1'

    row_data = [
        domain, company, web, vertical, 'Shopify' if is_shop else 'Khác',
        shop_id, rev_month, bw_rev, bw_spend, join_url, notes, 'Live (DNS OK)' if live else 'Không phản hồi'
    ]
    ws.append(row_data)

    if is_shop:
        ws.cell(row=r_idx, column=5).fill = shopify_fill
        ws.cell(row=r_idx, column=5).font = shopify_font

    if rev_month is not None:
        c_rev = ws.cell(row=r_idx, column=7)
        c_rev.number_format = '"$"#,##0'
        c_rev.alignment = Alignment(horizontal='right')

# Column widths
ws.column_dimensions['A'].width = 24
ws.column_dimensions['B'].width = 28
ws.column_dimensions['C'].width = 24
ws.column_dimensions['D'].width = 22
ws.column_dimensions['E'].width = 14
ws.column_dimensions['F'].width = 16
ws.column_dimensions['G'].width = 22
ws.column_dimensions['H'].width = 22
ws.column_dimensions['I'].width = 22
ws.column_dimensions['J'].width = 30
ws.column_dimensions['K'].width = 45
ws.column_dimensions['L'].width = 16

# Freeze header
ws.freeze_panes = 'A2'

wb.save(r'G:\AFF\Superfiliate_Enriched.xlsx')
print('Successfully created G:\\AFF\\Superfiliate_Enriched.xlsx')
