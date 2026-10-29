import sys
sys.stdout.reconfigure(encoding='utf-8')
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print('Reading Growi_Enriched.csv...')
with open(r'G:\AFF\Growi_Enriched.csv', 'r', encoding='utf-8', errors='replace') as f:
    rows = list(csv.reader(f))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Growi Deals (3991)'

# Styling
header_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')

shopify_fill = PatternFill(start_color='DCFCE7', end_color='DCFCE7', fill_type='solid')
shopify_font = Font(name='Segoe UI', size=10, color='166534', bold=True)

data_font = Font(name='Segoe UI', size=10)
num_font = Font(name='Segoe UI', size=10)

headers = [
    'Slug / Brand ID', 'Tên chương trình / Brand', 'Website (Domain)', '% Hoa hồng',
    'Ngành hàng', 'Shopify?', 'Shop ID', 'DT tháng (USD)', 'Link tham gia (Growi)', 'Ghi chú', 'Nguồn domain'
]

ws.append(headers)

for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for r_idx, r in enumerate(rows[1:], start=2):
    slug = r[0]
    prog_name = r[1]
    web = r[2]
    pct = float(r[3]) if r[3] else None
    cat = r[4]
    is_shop = r[5] == '1'
    shop_id = r[6]
    rev_month = float(r[7]) if r[7] else None
    join_url = r[8]
    notes = r[9]
    source = r[10]

    row_data = [
        slug, prog_name, web, pct, cat, 'Shopify' if is_shop else ('Có web' if web else 'TikTok Shop'),
        shop_id, rev_month, join_url, notes, source
    ]
    ws.append(row_data)

    # Style row
    if is_shop:
        ws.cell(row=r_idx, column=6).fill = shopify_fill
        ws.cell(row=r_idx, column=6).font = shopify_font

    if pct is not None:
        c_pct = ws.cell(row=r_idx, column=4)
        c_pct.number_format = '0.0"%"'
        c_pct.alignment = Alignment(horizontal='right')

    if rev_month is not None:
        c_rev = ws.cell(row=r_idx, column=8)
        c_rev.number_format = '"$"#,##0'
        c_rev.alignment = Alignment(horizontal='right')

# Auto fit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col[:100]:  # check first 100 rows for speed
        val_str = str(cell.value or '')
        if len(val_str) > max_len:
            max_len = len(val_str)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 26
ws.column_dimensions['C'].width = 26
ws.column_dimensions['D'].width = 14
ws.column_dimensions['E'].width = 16
ws.column_dimensions['F'].width = 14
ws.column_dimensions['G'].width = 16
ws.column_dimensions['H'].width = 18
ws.column_dimensions['I'].width = 38
ws.column_dimensions['J'].width = 40
ws.column_dimensions['K'].width = 16

# Freeze header
ws.freeze_panes = 'A2'

wb.save(r'G:\AFF\Growi_Enriched.xlsx')
print('Successfully created G:\\AFF\\Growi_Enriched.xlsx')
